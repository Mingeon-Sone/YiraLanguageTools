import os

import fitz

from docx import Document

import re



def search_and_read_pdf(file_path, search_word):

    try:

        with fitz.open(file_path) as pdf_doc:

            for page_num in range(pdf_doc.page_count):

                page = pdf_doc[page_num]

                text = page.get_text()

                if search_word.lower() in text.lower():

                    return True

        return False

    except Exception as e:

        print(f"Error reading {file_path} (pdf): {e}")

        return False



def extract_text_from_pdf(file_path):

    try:

        with fitz.open(file_path) as pdf_doc:

            text = ""

            for page_num in range(pdf_doc.page_count):

                page = pdf_doc[page_num]

                text += page.get_text()



            return text

    except Exception as e:

        print(f"Error extracting text from {file_path}: {e}")

        return ""



def search_and_read_docx(file_path, search_word, search_index=True):

    try:

        doc = Document(file_path)



        if not search_index:

            for paragraph in doc.paragraphs:

                if search_word.lower() in paragraph.text.lower():

                    return True

        else:

            index_text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])

            return search_word.lower() in index_text.lower()



        return False

    except Exception as e:

        print(f"Error reading {file_path}: {e}")

        return False



def search_in_file(file_path, search_word, search_index=True):

    if file_path.endswith(".docx"):

        return search_and_read_docx(file_path, search_word, search_index)

    elif file_path.endswith(".pdf"):

        return search_and_read_pdf(file_path, search_word)

    else:

        return False



def remove_enters(text):

    return text.replace('\n', ' ')



def slice_and_save_as_docx(file_path, output_folder, output_file_name, total_slices):

    try:

        pdf_text = extract_text_from_pdf(file_path)

        lines = pdf_text.split('\n')



        # Define a regular expression pattern to identify the question numbers

        question_pattern = re.compile(r'^\d+\. ')

        current_slice = 1

        doc_piece = Document()



        for line in lines:

            # Check if the line matches the question pattern

            if question_pattern.match(line.strip()):

                # Save the current piece and start a new one

                if current_slice <= total_slices:

                    output_file_path = os.path.join(output_folder, f"{output_file_name}_piece_{current_slice}.docx")

                    doc_piece.save(output_file_path)

                    print(f"Piece {current_slice} saved to: {output_file_path}")



                    current_slice += 1

                    doc_piece = Document()



            # Add the line to the current piece

            doc_piece.add_paragraph(line)



        # Save the last piece if necessary

        if current_slice <= total_slices:

            output_file_path = os.path.join(output_folder, f"{output_file_name}_piece_{current_slice}.docx")

            doc_piece.save(output_file_path)

            print(f"Piece {current_slice} saved to: {output_file_path}")



    except Exception as e:

        print(f"Error slicing {file_path}: {e}")



if __name__ == "__main__":

    folder_path = r'C:\Users\2rah\OneDrive\Desktop\project folder' # 탐색 폴더경로 입력

    results = []  # List to store found files



    while True:

        search_word = input("검색할 단어를 입력하세요 (나가려면 'exit'를 입력하세요): ")



        if search_word.lower() == 'exit':

            break



        search_index = input(folder_path + "(y/n): ").lower() == 'y'



        # Specify the desired output folder

        output_folder = r'C:\Users\2rah\OneDrive\Desktop\projectoutput' # 추출 결과 폴더경로 입력



        found = False



        for root, dirs, files in os.walk(folder_path):

            for file in files:

                if file.endswith((".docx", ".pdf")) and not file.startswith("~$"):

                    file_path = os.path.join(root, file)



                    if search_in_file(file_path, search_word, search_index):

                        print(f"Found '{search_word}' in: {file_path}")

                        results.append(file_path)



                        os.makedirs(output_folder, exist_ok=True)

                        output_file_name = f"{os.path.splitext(os.path.basename(file))[0]}_sliced"

                        slice_and_save_as_docx(file_path, output_folder, output_file_name, total_slices=45)



                        found = True



        if not found:

            print(f"'{search_word}' not found in any document.")

